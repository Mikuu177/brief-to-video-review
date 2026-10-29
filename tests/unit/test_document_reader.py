import copy
import io
import json
import zipfile
from pathlib import Path

import fitz
import pytest
from app.errors import ReviewError
from app.services.document_inventory import policy
from app.services.document_reader import archive_inventory, read_document, read_txt, text_agreement
from PIL import Image, ImageDraw


def test_short_prohibition_is_not_blank(tmp_path):
    path = tmp_path / "short.pdf"
    with fitz.open() as doc:
        page = doc.new_page(width=400, height=300)
        page.insert_text((40, 90), "DO NOT CLAIM", fontsize=20)
        doc.save(path)
    result = read_document(path, tmp_path / "out", policy())
    page = result["units"][0]
    assert len(page["native_text"].strip()) < 30
    assert page["features"]["blank_candidate"] is False
    assert "DO NOT CLAIM" in page["native_text"]
    assert "CLAIM" in page["ocr_text"].upper()
    assert result["inventory"]["independent_page_count"] == 1


def test_scanned_page_has_independent_raster_text(tmp_path):
    image = Image.new("RGB", (1000, 600), "white")
    draw = ImageDraw.Draw(image)
    # Scaled text yields a controlled raster-only source; it is not a human label.
    draw.text((30, 40), "NOT ALLOWED", fill="black", font_size=90)
    data = io.BytesIO()
    image.save(data, "PNG")
    path = tmp_path / "scan.pdf"
    with fitz.open() as doc:
        page = doc.new_page(width=500, height=300)
        page.insert_image(page.rect, stream=data.getvalue())
        doc.save(path)
    result = read_document(path, tmp_path / "out", policy())
    page = result["units"][0]
    assert not page["native_text"].strip()
    assert "ALLOWED" in page["ocr_text"].upper()
    assert page["state"] == "NEEDS_HUMAN"
    assert any(u["kind"] == "image" for u in result["units"])


def test_page_budget_blocks_instead_of_truncating(tmp_path):
    path = tmp_path / "three.pdf"
    with fitz.open() as doc:
        for _ in range(3):
            doc.new_page()
        doc.save(path)
    limited = copy.deepcopy(policy())
    limited["p0"]["max_pdf_pages"] = 2
    with pytest.raises(ReviewError, match="未只读前200页"):
        read_document(path, tmp_path / "out", limited)
    assert not list((tmp_path / "out").glob("*.png"))


def test_txt_keeps_last_requirement_and_blank_lines(tmp_path):
    path = tmp_path / "source.txt"
    path.write_text("背景\n\n不得说立刻清醒。\n")
    result = read_txt(path, policy())
    assert len(result["units"]) == 3
    assert result["units"][1]["native_text"] == ""
    assert result["units"][-1]["native_text"] == "不得说立刻清醒。"
    assert result["units"][-1]["locator"]["line_start"] == 3


def test_text_limit_does_not_return_partial_result(tmp_path):
    path = tmp_path / "long.txt"
    path.write_text("不" * 101)
    limited = copy.deepcopy(policy())
    limited["reader_limits"]["max_text_characters"] = 100
    with pytest.raises(ReviewError) as exc:
        read_txt(path, limited)
    assert exc.value.code == "TEXT_CHARACTER_BUDGET"


def test_empty_txt_is_not_a_complete_source(tmp_path):
    path = tmp_path / "empty.txt"
    path.write_text("\n  \n")
    with pytest.raises(ReviewError) as exc:
        read_txt(path, policy())
    assert exc.value.code == "EMPTY_TEXT_SOURCE"


def test_zip_expansion_is_bounded(tmp_path):
    path = tmp_path / "large.docx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * 200000)
    with pytest.raises(ReviewError) as exc:
        archive_inventory(path, policy()["reader_limits"])
    assert exc.value.code == "ARCHIVE_SIZE_BUDGET"


def test_docx_registers_auxiliary_parts_and_revisions(tmp_path):
    from docx import Document

    path = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("正文的要求")
    doc.add_table(rows=1, cols=2).cell(0, 1).text = "不得改产品名"
    doc.sections[0].header.paragraphs[0].text = "页眉禁用要求"
    doc.save(path)
    # Introduce a real revision XML part, then read without silently accepting it.
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["word/document.xml"] = parts["word/document.xml"].replace(
        b"</w:body>", b'<w:ins w:id="1"><w:r><w:t>inserted condition</w:t></w:r></w:ins></w:body>'
    )
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    result = read_document(path, tmp_path / "out", policy())
    assert "页眉禁用要求" in "".join(u["native_text"] for u in result["units"])
    assert any(u["kind"] == "table_cell" and "不得" in u["native_text"] for u in result["units"])
    codes = {i["code"] for u in result["units"] for i in u["issues"]}
    assert "REVISION_POLICY_REQUIRED" in codes
    assert "DOCX_RENDER_NOT_QUALIFIED" in codes
    assert result["inventory"]["physical_pages"] is None


def test_legacy_doc_remains_unqualified(tmp_path):
    path = tmp_path / "legacy.doc"
    path.write_bytes(b"legacy")
    with pytest.raises(ReviewError) as exc:
        read_document(path, tmp_path / "out", policy())
    assert exc.value.code == "FORMAT_NOT_QUALIFIED"


def test_ocr_agreement_is_a_diagnostic_not_a_semantic_score():
    assert text_agreement("不得提神", "不得提神") == 1
    assert text_agreement("不得提神", "可以提神") < 1
    assert text_agreement("", "图中有字") is None


def test_policy_and_gold_protocol_do_not_claim_human_truth():
    root = Path(__file__).resolve().parents[2]
    protocol = json.loads((root / "evaluation/gold_standard/brief_protocol_v1.json").read_text())
    assert protocol["confirmed_labels"] == 0 and protocol["metrics"] is None
    assert protocol["sources"] == []  # No private source metadata is published.


def test_pdf_count_disagreement_is_not_success(tmp_path, monkeypatch):
    document = fitz.open()
    document.new_page()
    path = tmp_path / "mismatch.pdf"
    document.save(path)
    monkeypatch.setattr("app.services.document_reader.pdf_count", lambda path: 2)
    with pytest.raises(ReviewError) as exc:
        read_document(path, tmp_path / "output", policy())
    assert exc.value.code == "PDF_PAGE_COUNT_MISMATCH"


def test_rotation_annotations_and_embedded_attachment_are_accounted(tmp_path):
    document = fitz.open()
    page = document.new_page(width=500, height=200)
    page.insert_text((350, 50), "NO PROMISE", fontsize=12)
    page.add_text_annot((20, 30), "Separate mandatory comment")
    page.insert_link(
        {"kind": fitz.LINK_URI, "from": fitz.Rect(10, 10, 30, 30), "uri": "https://example.invalid/appendix"}
    )
    page.set_rotation(90)
    document.embfile_add("additional-brief.txt", b"Additional requirement")
    path = tmp_path / "rotated.pdf"
    document.save(path)
    result = read_document(path, tmp_path / "output", policy())
    units = result["units"]
    assert {"page:1", "attachment:0"} == set(result["inventory"]["expected_root_keys"])
    assert next(u for u in units if u["kind"] == "attachment")["state"] == "BLOCKED"
    assert next(u for u in units if u["kind"] == "annotation")["native_text"] == "Separate mandatory comment"
    assert next(u for u in units if u["kind"] == "external_reference")["state"] == "BLOCKED"
    native = next(u for u in units if u["kind"] == "text_region" and "NO PROMISE" in u["native_text"])
    assert all(0 <= coordinate <= 1 for coordinate in native["locator"]["bbox"])
    image = Image.open(tmp_path / "output/page-1.png").convert("L")
    x0, y0, x1, y1 = native["locator"]["bbox"]
    crop = image.crop(
        (int(x0 * image.width), int(y0 * image.height), int(x1 * image.width), int(y1 * image.height))
    )
    assert sum(crop.histogram()[:128]) > 10


def test_visible_page_with_no_recognized_text_is_blocked(tmp_path, monkeypatch):
    document = fitz.open()
    page = document.new_page()
    page.draw_rect(fitz.Rect(40, 40, 300, 300), color=(0, 0, 0), fill=(0, 0, 0))
    path = tmp_path / "not-blank.pdf"
    document.save(path)
    monkeypatch.setattr(
        "app.services.document_reader.ocr_image", lambda *a, **k: {"text": "", "paragraphs": []}
    )
    root = read_document(path, tmp_path / "output", policy())["units"][0]
    assert root["state"] == "BLOCKED"
    assert root["features"]["blank_candidate"] is False


def test_docx_table_dimensions_fail_instead_of_truncating(tmp_path):
    path = tmp_path / "table.docx"
    with zipfile.ZipFile(path, "w") as package:
        package.writestr(
            "word/document.xml",
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:tbl><w:tr><w:tc><w:p><w:r><w:t>A</w:t></w:r></w:p></w:tc></w:tr><w:tr><w:tc><w:p><w:r><w:t>last mandatory cell</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:body></w:document>',
        )
    limited = copy.deepcopy(policy())
    limited["reader_limits"]["max_table_rows"] = 1
    with pytest.raises(ReviewError) as exc:
        read_document(path, tmp_path / "output", limited)
    assert exc.value.code == "TABLE_DIMENSION_BUDGET"
