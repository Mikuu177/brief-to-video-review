import hashlib
import io
import json

import pymupdf
import pytest
from app.errors import ReviewError
from app.parsers.documents import docx, legacy_doc, pdf, txt
from conftest import upload_text
from docx import Document
from PIL import Image


def test_corrupt_legacy_doc_is_not_a_success(tmp_path):
    path = tmp_path / "broken.doc"
    path.write_bytes(bytes.fromhex("d0cf11e0a1b11ae1") + b"broken document")
    with pytest.raises(ReviewError, match="提取失败"):
        legacy_doc(path, "test-broken")



def test_docx_interleaved_tables_keep_order_and_txt_keeps_lines(tmp_path):
    doc = Document()
    doc.add_paragraph("第一段")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "画面"
    table.cell(0, 1).text = "台词"
    doc.add_paragraph("第二段")
    path = tmp_path / "test.docx"
    doc.save(path)
    blocks, _, _ = docx(path, "test")
    assert [b["text"] for b in blocks] == ["第一段", "画面", "台词", "第二段"]
    path = tmp_path / "test.txt"
    path.write_bytes("第一行\n\n第三行".encode("gb18030"))
    blocks, _, _ = txt(path, "test")
    assert blocks[-1]["locator"]["line_start"] == 3


def test_scan_blank_and_bad_pdf_are_distinct(tmp_path):
    document = pymupdf.open()
    document.new_page()
    page = document.new_page()
    image = Image.new("RGB", (200, 100), "white")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    page.insert_image(page.rect, stream=buffer.getvalue())
    path = tmp_path / "scan.pdf"
    document.save(path)
    _, coverage, _ = pdf(path, "test")
    assert coverage[0]["status"] == "blank"
    assert coverage[1]["status"] == "needs_ocr"


def test_file_signature_path_and_duplicate_hash(environment, project, monkeypatch):
    client, _ = environment
    p = project["id"]
    r = client.post(
        f"/api/v1/projects/{p}/assets",
        data={"kind": "brief"},
        files={"file": ("forged.pdf", b"not a pdf", "application/pdf")},
    )
    assert r.status_code == 422 and r.json()["code"] == "INVALID_FILE_TYPE"
    r = client.post(
        f"/api/v1/projects/{p}/assets",
        data={"kind": "brief"},
        files={"file": ("../../escape.txt", "原文内容".encode(), "text/plain")},
    )
    assert r.status_code == 201 and r.json()["filename"] == "escape.txt"
    first = r.json()
    r = client.post(
        f"/api/v1/projects/{p}/assets",
        data={"kind": "brief"},
        files={"file": ("another.txt", "原文内容".encode(), "text/plain")},
    )
    assert r.json()["sha256"] == first["sha256"] and r.json()["id"] != first["id"]
    assert first["sha256"] == hashlib.sha256("原文内容".encode()).hexdigest()
    from app.settings import settings

    monkeypatch.setattr(settings, "review_max_bytes", 10)
    response = client.post(
        f"/api/v1/projects/{p}/assets",
        data={"kind": "brief"},
        files={"file": ("large.txt", b"elevenbytes!", "text/plain")},
    )
    assert response.status_code == 422 and response.json()["code"] == "FILE_TOO_LARGE"


def test_video_range_download_and_unauthorized_access(environment, project, sample_video):
    client, _ = environment
    path = sample_video
    response = client.post(
        f"/api/v1/projects/{project['id']}/assets",
        data={"kind": "video"},
        files={"file": (path.name, path.read_bytes(), "video/mp4")},
    )
    assert response.status_code == 201
    url = f"/api/v1/assets/{response.json()['id']}/content"
    partial = client.get(url, headers={"Range": "bytes=100-199"})
    assert partial.status_code == 206 and len(partial.content) == 100
    assert partial.content == path.read_bytes()[100:200]
    client.post("/api/v1/sessions", json={"username": "test-reviewer", "password": "review-password"})
    assert client.get(url).status_code == 403
    assert client.get(f"/api/v1/projects/{project['id']}").status_code == 403


def test_cross_origin_mutation_and_secret_exposure(environment, project):
    client, _ = environment
    r = client.post("/api/v1/projects", headers={"Origin": "https://untrusted.example"}, json={})
    assert r.status_code == 403 and r.json()["code"] == "INVALID_ORIGIN"
    config = client.get("/api/v1/config").json()
    assert not config["text"]["ready"] and not config["video"]["ready"]
    assert "ark_api_key" not in json.dumps(config) and "password" not in json.dumps(config)


def test_material_change_cannot_overwrite_existing_source(environment, project):
    client, _ = environment
    source = upload_text(client, project["id"], content="第一版原文")
    old_id = source["asset"]["id"]
    upload_text(client, project["id"], content="第二版原文")
    assert client.get(f"/api/v1/assets/{old_id}/content").content == "第一版原文".encode()


def test_synthetic_27_page_pdf_is_fully_accounted_for(tmp_path, monkeypatch):
    from app.settings import settings

    monkeypatch.setattr(settings, "data_dir", tmp_path)
    document = pymupdf.open()
    for number in range(1, 28):
        page = document.new_page()
        page.insert_text((60, 60), f"Synthetic Brief page {number}")
        if number == 20:
            page.insert_text((60, 100), "Product must appear within 15 seconds.")
    path = tmp_path / "synthetic-brief.pdf"
    document.save(path)
    document.close()
    blocks, coverage, metadata = pdf(path, "synthetic-pdf")
    assert metadata["page_count"] == len(coverage) == 27
    assert {p["page"] for p in coverage} == set(range(1, 28))
    assert any("15 seconds" in b["text"] and b["locator"]["page"] == 20 for b in blocks)
