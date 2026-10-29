import json
import posixpath
import re
import subprocess
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import fitz
from app.errors import ReviewError
from app.services.storage import decode_text, local_path
from app.settings import ROOT
from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook


def block(text: str, locator: dict, method: str, status="success") -> dict:
    return {"text": text, "locator": locator, "extraction_method": method, "coverage_status": status}


def pdf(path: Path, asset_id: str) -> tuple[list, list, dict]:
    blocks, coverage = [], []
    with fitz.open(path) as document:
        if document.is_encrypted:
            raise ReviewError("PARSE_FAILED", "PDF 受密码保护")
        for index, page in enumerate(document):
            texts = []
            for item in page.get_text("blocks"):
                if item[6] == 0 and item[4].strip():
                    texts.append(item[4].strip())
                    bbox = [
                        item[0] / page.rect.width,
                        item[1] / page.rect.height,
                        item[2] / page.rect.width,
                        item[3] / page.rect.height,
                    ]
                    blocks.append(
                        block(
                            item[4].strip(),
                            {"kind": "pdf", "page": index + 1, "bbox": bbox},
                            "pymupdf_native",
                        )
                    )
            has_visual = bool(page.get_images() or page.get_drawings())
            native_chars = len("".join(texts))
            status = "success" if native_chars >= 30 else ("needs_ocr" if has_visual else "blank")
            # Native text is not proof that infographic content was fully parsed.
            coverage.append(
                {
                    "page": index + 1,
                    "status": status,
                    "native_characters": native_chars,
                    "has_visual": has_visual,
                    "visual_review": "pending" if has_visual else "not_required",
                }
            )
        return (
            blocks,
            coverage,
            {
                "page_count": len(document),
                "limits": ["原生文字已提取；图片、版面分组及图中要求需视觉解析或人工核对。"],
            },
        )


def txt(path: Path, asset_id: str) -> tuple[list, list, dict]:
    text = decode_text(path.read_bytes())
    lines = text.splitlines()
    blocks = [
        block(line, {"kind": "txt", "line_start": i + 1, "line_end": i + 1}, "text_decode")
        for i, line in enumerate(lines)
        if line.strip()
    ]
    if not blocks:
        raise ReviewError("PARSE_FAILED", "没有可读取的文本")
    return blocks, [{"status": "success", "lines": len(lines)}], {}


def docx(path: Path, asset_id: str) -> tuple[list, list, dict]:
    document = Document(path)
    blocks = []
    paragraph_index, table_index = 0, 0
    # Walk document body so interleaved paragraphs and tables retain their original order.
    for child in document.element.body:
        if child.tag == qn("w:p"):
            text = "".join(node.text or "" for node in child.iter(qn("w:t")))
            if text.strip():
                blocks.append(
                    block(text, {"kind": "docx", "paragraph_index": paragraph_index}, "python_docx")
                )
            paragraph_index += 1
        elif child.tag == qn("w:tbl"):
            seen = set()
            table = document.tables[table_index]
            for row_index, row in enumerate(table.rows):
                for col_index, cell in enumerate(row.cells):
                    if cell._tc in seen:
                        continue
                    seen.add(cell._tc)
                    if cell.text.strip():
                        blocks.append(
                            block(
                                cell.text,
                                {
                                    "kind": "docx",
                                    "table_index": table_index,
                                    "row": row_index,
                                    "column": col_index,
                                },
                                "python_docx",
                            )
                        )
            table_index += 1
    if not blocks:
        raise ReviewError("PARSE_FAILED", "Word 未提取到正文")
    return (
        blocks,
        [{"status": "partial", "reason": "正文与表格已解析；图片和修订需核对"}],
        {"limits": ["Word 不提供稳定物理页码；图片、批注与修订尚未完整解析。"]},
    )


def legacy_doc(path: Path, asset_id: str) -> tuple[list, list, dict]:
    try:
        result = subprocess.run(
            ["node", str(ROOT / "scripts/extract_doc.cjs"), str(path)],
            capture_output=True,
            timeout=60,
            check=True,
        )
        text = json.loads(result.stdout)["body"].strip()
    except (subprocess.SubprocessError, OSError, ValueError, KeyError) as exc:
        raise ReviewError("PARSE_FAILED", "旧 DOC 正文提取失败，请转换为 DOCX 后重试") from exc
    if not text:
        raise ReviewError("PARSE_FAILED", "旧 DOC 提取结果为空")
    blocks = [
        block(
            line,
            {"kind": "legacy_doc_text", "paragraph_index": index},
            "word_extractor_text_fallback",
            "partial",
        )
        for index, line in enumerate(text.splitlines())
        if line.strip()
    ]
    return (
        blocks,
        [{"status": "partial", "reason": "仅旧 DOC 正文兜底提取"}],
        {"limits": ["来自旧 DOC 的文本兜底，不是 DOCX 转换件；表格布局、图片、修订和批注未完整覆盖。"]},
    )


def xlsx(path: Path, asset_id: str) -> tuple[list, list, dict]:
    images = {}
    with zipfile.ZipFile(path) as archive:
        if "xl/cellimages.xml" in archive.namelist():
            ns = {
                "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
                "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
                "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
            }
            rels = {
                node.attrib["Id"]: posixpath.normpath("xl/" + node.attrib["Target"])
                for node in ET.fromstring(archive.read("xl/_rels/cellimages.xml.rels"))
            }
            for picture in ET.fromstring(archive.read("xl/cellimages.xml")).findall(".//xdr:pic", ns):
                name = picture.find(".//xdr:cNvPr", ns).attrib["name"]
                rid = picture.find(".//a:blip", ns).attrib[f"{{{ns['r']}}}embed"]
                source = rels[rid]
                key = f"previews/{asset_id}/{Path(source).name}"
                dest = local_path(key)
                dest.parent.mkdir(parents=True, exist_ok=True)
                dest.write_bytes(archive.read(source))
                images[name] = key
    workbook = load_workbook(path, data_only=False)
    blocks, coverage = [], []
    for sheet in workbook:
        count = 0
        for row in sheet:
            for cell in row:
                if cell.value is None:
                    continue
                value = str(cell.value)
                locator = {"kind": "xlsx", "sheet": sheet.title, "cell": cell.coordinate}
                for merged in sheet.merged_cells.ranges:
                    if cell.coordinate in merged:
                        locator["merged_range"] = str(merged)
                if "DISPIMG" in value:
                    match = re.search(r'"(ID_[^"]+)"', value)
                    if match and match.group(1) in images:
                        locator["image_key"] = images[match.group(1)]
                        value = f"[内嵌图片：{cell.coordinate}，内容待视觉审核]"
                blocks.append(block(value, locator, "openpyxl_wps"))
                count += 1
        coverage.append(
            {
                "sheet": sheet.title,
                "status": "partial" if images else "success",
                "nonempty_cells": count,
                "merged_ranges": [str(x) for x in sheet.merged_cells.ranges],
            }
        )
    if not blocks:
        raise ReviewError("PARSE_FAILED", "工作簿没有可读取的单元格")
    return (
        blocks,
        coverage,
        {
            "embedded_images": len(images),
            "limits": ["内嵌图片已关联单元格；图像内容需模型或人工检查。"] if images else [],
        },
    )


PARSERS = {".pdf": pdf, ".txt": txt, ".docx": docx, ".doc": legacy_doc, ".xlsx": xlsx}


def parse_document(path: Path, asset_id: str):
    parser = PARSERS.get(path.suffix.lower())
    if not parser:
        raise ReviewError("UNSUPPORTED_FILE_TYPE", "该文件不能作为文档解析")
    return parser(path, asset_id)
